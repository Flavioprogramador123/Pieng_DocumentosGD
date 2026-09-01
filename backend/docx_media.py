"""
Inserção de imagens inline em DOCX (memorial e anexos).
"""

from __future__ import annotations

from pathlib import Path


def _paragraph_plain_text(paragraph) -> str:
    return paragraph.text or ''


def _clear_paragraph_runs(paragraph) -> None:
    """Remove conteúdo do parágrafo preservando w:pPr (estilo/alinhamento)."""
    from docx.oxml.ns import qn

    p_el = paragraph._element
    for child in list(p_el):
        if child.tag != qn('w:pPr'):
            p_el.remove(child)


def _iter_all_paragraphs(doc):
    for paragraph in doc.paragraphs:
        yield paragraph
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
    for section in doc.sections:
        for part in (
            section.header,
            section.footer,
            getattr(section, 'first_page_header', None),
            getattr(section, 'first_page_footer', None),
        ):
            if part is None:
                continue
            for paragraph in part.paragraphs:
                yield paragraph
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            yield paragraph


def insert_image_at_placeholders(
    docx_path: Path,
    image_path: Path,
    placeholders: tuple[str, ...],
    *,
    width_cm: float = 12.0,
) -> bool:
    if not image_path.is_file():
        return False
    try:
        from docx import Document
        from docx.shared import Cm
    except ImportError:
        return False

    doc = Document(str(docx_path))

    def replace_in_paragraph(paragraph) -> bool:
        text = _paragraph_plain_text(paragraph)
        if not any(p in text for p in placeholders):
            return False
        _clear_paragraph_runs(paragraph)
        run = paragraph.add_run()
        run.add_picture(str(image_path), width=Cm(width_cm))
        return True

    for paragraph in _iter_all_paragraphs(doc):
        if replace_in_paragraph(paragraph):
            doc.save(str(docx_path))
            return True

    return False
