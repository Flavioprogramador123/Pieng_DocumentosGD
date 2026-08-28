import pandas as pd
from openpyxl import load_workbook

def fill_excel_form(data, template_path, output_path):
    wb = load_workbook(template_path)
    ws = wb["0"]  # Assuming '0' is the main sheet for data entry

    # Example: Filling client personal data (adjust cell references as needed)
    ws["C15"] = data["client_name"]
    ws["C16"] = data["client_address"]
    ws["C17"] = data["consumer_unit"]

    # Example: Filling technical system data (adjust cell references as needed)
    # Assuming '1' is the sheet for modules/sources
    ws_sources = wb["1"]
    for i, module in enumerate(data["modules"]):
        row = 10 + i  # Adjust starting row as needed
        ws_sources.cell(row=row, column=2, value=module["quantity"])
        ws_sources.cell(row=row, column=3, value=module["model"])
        ws_sources.cell(row=row, column=4, value=module["power"])

    # Example: Filling inverter data (adjust cell references as needed)
    # Assuming '2' is the sheet for inverters
    ws_inverters = wb["2"]
    for i, inverter in enumerate(data["inverters"]):
        row = 10 + i  # Adjust starting row as needed
        ws_inverters.cell(row=row, column=2, value=inverter["quantity"])
        ws_inverters.cell(row=row, column=3, value=inverter["model"])
        ws_inverters.cell(row=row, column=4, value=inverter["power"])

    # Save the filled workbook
    wb.save(output_path)

    print(f"Excel form filled and saved to {output_path}")

if __name__ == "__main__":
    # Sample data (this would come from user input or another source)
    sample_data = {
        "client_name": "João da Silva",
        "client_address": "Rua das Flores, 123, Bairro Centro, Cidade, Estado, CEP",
        "consumer_unit": "1234567-8",
        "modules": [
            {"quantity": 10, "model": "ModuleX-350W", "power": 350},
            {"quantity": 5, "model": "ModuleY-400W", "power": 400},
        ],
        "inverters": [
            {"quantity": 1, "model": "InverterA-5kW", "power": 5000},
        ],
    }

    template_excel_path = "/home/ubuntu/equatorial_automation/templates/NT.00020-05-Anexo-I-Formulario-de-Solicitacao-de-Orcamento-de-Microgeracao-Distribuida-Grupo-B.xlsx"
    output_excel_path = "/home/ubuntu/equatorial_automation/output/Formulario_Orcamento_Preenchido.xlsx"

    fill_excel_form(sample_data, template_excel_path, output_excel_path)


