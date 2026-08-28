from pathlib import Path
from openpyxl import load_workbook

base = Path('/home/ubuntu/work_equatorial')
for path in sorted(base.glob('*.xltx')):
    print(f'=== {path.name} ===')
    wb = load_workbook(path, read_only=False, data_only=False)
    print('SHEETS:', wb.sheetnames)
    for ws in wb.worksheets:
        print(f'-- {ws.title} ({ws.max_row} rows x {ws.max_column} cols) --')
        for row in ws.iter_rows():
            values = []
            for cell in row:
                if cell.value is not None:
                    values.append(f'{cell.coordinate}={cell.value!r}')
            if values:
                print(' | '.join(values))
    print()
