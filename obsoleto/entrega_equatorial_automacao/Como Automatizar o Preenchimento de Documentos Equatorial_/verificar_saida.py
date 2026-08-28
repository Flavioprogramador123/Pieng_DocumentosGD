from pathlib import Path
from zipfile import ZipFile, BadZipFile
from openpyxl import load_workbook
import re

base = Path('/home/ubuntu/work_equatorial/saida_exemplo')
for path in sorted(base.iterdir()):
    if path.suffix.lower() not in {'.docx', '.xlsx'}:
        continue
    print(f'=== {path.name} ===')
    try:
        with ZipFile(path) as archive:
            bad = archive.testzip()
            print('ZIP:', 'OK' if bad is None else f'ERRO em {bad}')
            all_xml = ''
            for name in archive.namelist():
                if name.endswith('.xml'):
                    all_xml += archive.read(name).decode('utf-8', errors='ignore') + '\n'
            tokens = sorted(set(re.findall(r'\{\{[^}]+\}\}', all_xml)))
            print('TOKENS:', tokens)
            print('HAS_LOGO:', 'word/media/image1.png' in archive.namelist())
    except BadZipFile as exc:
        print('ZIP: ERRO', exc)
    if path.suffix.lower() == '.xlsx':
        wb = load_workbook(path, data_only=False, read_only=False)
        ws = wb['1']
        for coordinate in ['C10', 'R10', 'AC10', 'C13', 'D15', 'I15', 'Q15', 'V15', 'Z17', 'P33', 'Y33', 'C38', 'G49', 'G51']:
            print(f'{coordinate}={ws[coordinate].value!r}')
