from pathlib import Path
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

OUT = Path('/home/ubuntu/work_equatorial/tabela_demanda_industrial_20kw.docx')

rows = [
    ('1', 'Iluminação industrial e áreas de circulação', '500', '4', '2,00', '0,95', '2,11', '100%', '2,00', '2,11'),
    ('2', 'Tomadas de uso geral e equipamentos auxiliares', '750', '4', '3,00', '0,90', '3,33', '80%', '2,40', '2,67'),
    ('3', 'Compressor de ar industrial', '7.500', '1', '7,50', '0,85', '8,82', '80%', '6,00', '7,06'),
    ('4', 'Bomba/motor de processo', '5.500', '1', '5,50', '0,85', '6,47', '80%', '4,40', '5,18'),
    ('5', 'Máquina de solda / equipamento de produção', '5.000', '1', '5,00', '0,80', '6,25', '80%', '4,00', '5,00'),
    ('6', 'Escritório, informática e apoio', '2.000', '1', '2,00', '0,95', '2,11', '60%', '1,20', '1,27'),
]

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('LEVANTAMENTO DE CARGA – UNIDADE INDUSTRIAL')
r.bold = True
r.font.size = Pt(14)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Tabela de referência com demanda-alvo de aproximadamente 20 kW')
r.italic = True
r.font.size = Pt(10)

p = doc.add_paragraph('Cliente: ROSEMBERGUE LEAO E SILVA | UC: 000068889901235 | Classe: Industrial')
p.paragraph_format.space_after = Pt(6)

headers = ['Item', 'Descrição', 'Pot. Unit. (W)', 'Qtd.', 'CI (kW)', 'FP', 'CI (kVA)', 'FD', 'D (kW)', 'D (kVA)']
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
for cell, text in zip(table.rows[0].cells, headers):
    cell.text = text
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(8)

for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if cell is not cells[1] else WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                run.font.size = Pt(8)

cells = table.add_row().cells
summary = ['TOTAL', 'Demanda de projeto a confirmar pelo técnico', '—', '—', '25,00', '—', '29,09', '—', '20,00', '23,29']
for cell, text in zip(cells, summary):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if cell is not cells[1] else WD_ALIGN_PARAGRAPH.LEFT
        for run in paragraph.runs:
            run.bold = True
            run.font.size = Pt(8)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(8)
p.add_run('Observação técnica: ').bold = True
p.add_run('esta tabela é um rascunho de referência solicitado para uma demanda-alvo de 20 kW. Os equipamentos, potências unitárias, fatores de potência, fatores de demanda e simultaneidade devem ser conferidos pelo técnico responsável e substituídos pelos valores reais da instalação antes do protocolo.')

p = doc.add_paragraph()
p.add_run('Atenção ao total: ').bold = True
p.add_run('a composição acima resulta em 20,00 kW de demanda calculada como referência. O técnico deve substituir as cargas, fatores de potência e fatores de demanda pelos valores reais da instalação; não é recomendável manter qualquer valor apenas para atingir o total solicitado.')

doc.save(OUT)
print(OUT)
