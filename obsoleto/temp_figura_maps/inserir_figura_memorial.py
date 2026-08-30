#!/usr/bin/env python3
"""
Protótipo (.temp) — tenta colar a PNG no memorial DOCX (cópia).

Substitui o parágrafo que contém {{FIGURA_LOCALIZACAO}} ou o texto placeholder
por uma imagem inline. Gera arquivo *_com_figura.docx ao lado; não sobrescreve o original.

Requer: pip install python-docx (só nesta pasta .temp)
"""

from __future__ import annotations

import argparse
import shutil
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Cm
except ImportError as exc:
    raise SystemExit(
        'Instale python-docx na venv .temp:\n'
        '  .venv\\Scripts\\pip install python-docx'
    ) from exc

PLACEHOLDERS = (
    '{{FIGURA_LOCALIZACAO}}',
    '[Inserir figura / print do mapa da localização]',
    'Figura tirado do maps da localização do imóvel',
)


def insert_image(docx_in: Path, png: Path, docx_out: Path, width_cm: float = 14.0) -> bool:
    doc = Document(docx_in)
    replaced = False

    for paragraph in doc.paragraphs:
        text = paragraph.text or ''
        if not any(p in text for p in PLACEHOLDERS):
            continue
        for run in paragraph.runs:
            run.text = ''
        paragraph.text = ''
        run = paragraph.add_run()
        run.add_picture(str(png), width=Cm(width_cm))
        replaced = True
        break

    if not replaced:
        # tabelas
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        text = paragraph.text or ''
                        if not any(p in text for p in PLACEHOLDERS):
                            continue
                        for run in paragraph.runs:
                            run.text = ''
                        paragraph.text = ''
                        run = paragraph.add_run()
                        run.add_picture(str(png), width=Cm(width_cm))
                        replaced = True
                        break
                    if replaced:
                        break
                if replaced:
                    break
            if replaced:
                break

    if not replaced:
        return False

    doc.save(docx_out)
    return True


def main() -> int:
    parser = argparse.ArgumentParser(description='Insere PNG no memorial (protótipo .temp)')
    parser.add_argument('--memorial', type=Path, required=True, help='DOCX memorial (será copiado)')
    parser.add_argument('--figura', type=Path, required=True, help='PNG gerada por gerar_figura_maps.py')
    parser.add_argument('-o', '--output', type=Path, help='DOCX de saída')
    parser.add_argument('--largura-cm', type=float, default=14.0)
    args = parser.parse_args()

    if not args.memorial.is_file():
        raise SystemExit(f'Memorial não encontrado: {args.memorial}')
    if not args.figura.is_file():
        raise SystemExit(f'Figura não encontrada: {args.figura}')

    out = args.output or args.memorial.with_name(args.memorial.stem + '_com_figura.docx')
    tmp = out.with_suffix('.tmp.docx')
    shutil.copy2(args.memorial, tmp)

    ok = insert_image(tmp, args.figura, out, width_cm=args.largura_cm)
    tmp.unlink(missing_ok=True)

    if not ok:
        raise SystemExit(
            'Placeholder {{FIGURA_LOCALIZACAO}} não encontrado no DOCX.\n'
            'Abra o memorial e confira se o token ainda está no texto.'
        )

    print(f'OK: {out}')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
