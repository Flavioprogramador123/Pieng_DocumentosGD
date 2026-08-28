from flask import Blueprint, request, jsonify
from flask_cors import cross_origin
import pandas as pd
from openpyxl import load_workbook
import os
import sys
from docx import Document
import json
from advanced_calculations import (
    calculate_cable_section_advanced,
    calculate_protection_devices_advanced,
    calculate_energy_generation,
    calculate_economic_analysis,
    validate_system_compatibility
)

automation_bp = Blueprint('automation', __name__)

@automation_bp.route('/fill-documents', methods=['POST'])
@cross_origin()
def fill_documents():
    try:
        data = request.json
        print(f"Received data: {data}")

        # Validate required fields
        required_fields = ['client_name', 'client_address', 'consumer_unit']
        for field in required_fields:
            if field not in data:
                return jsonify({'error': f'Missing required field: {field}'}), 400

        # Fill Excel form
        print("Filling Excel form...")
        excel_result = fill_excel_form(data)

        # Fill Word documents
        print("Filling memorial document...")
        memorial_result = fill_memorial_document(data)
        print("Filling procuracao document...")
        procuracao_result = fill_procuracao_document(data)

        return jsonify({
            'success': True,
            'files': {
                'excel': excel_result,
                'memorial': memorial_result,
                'procuracao': procuracao_result
            }
        })

    except Exception as e:
        import traceback
        error_trace = traceback.format_exc()
        print(f"ERROR in fill_documents: {error_trace}")
        return jsonify({'error': str(e), 'trace': error_trace}), 500

def fill_excel_form(data):
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
    template_path = os.path.join(base_dir, "templates", "NT.00020-05-Anexo-I-Formulario-de-Solicitacao-de-Orcamento-de-Microgeracao-Distribuida-Grupo-B.xlsx")
    output_path = os.path.join(base_dir, "output", "Formulario_Orcamento_Preenchido.xlsx")

    # Create output directory if it doesn't exist
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    # Create a copy of the template
    wb = load_workbook(template_path)
    
    # Fill basic client data (adjust cell references based on actual form structure)
    if '0' in wb.sheetnames:
        ws = wb['0']
        # These are example cell references - need to be adjusted based on actual form
        try:
            ws['C15'] = data.get('client_name', '')
            ws['C16'] = data.get('client_address', '')
            ws['C17'] = data.get('consumer_unit', '')
        except:
            pass  # Handle merged cells or protected cells
    
    # Fill modules data if provided
    if 'modules' in data and '1' in wb.sheetnames:
        ws_modules = wb['1']
        for i, module in enumerate(data['modules']):
            row = 15 + i  # Adjust starting row
            try:
                ws_modules.cell(row=row, column=2, value=module.get('quantity', ''))
                ws_modules.cell(row=row, column=3, value=module.get('model', ''))
                ws_modules.cell(row=row, column=4, value=module.get('power', ''))
            except:
                pass
    
    # Fill inverters data if provided
    if 'inverters' in data and '2' in wb.sheetnames:
        ws_inverters = wb['2']
        for i, inverter in enumerate(data['inverters']):
            row = 15 + i  # Adjust starting row
            try:
                ws_inverters.cell(row=row, column=2, value=inverter.get('quantity', ''))
                ws_inverters.cell(row=row, column=3, value=inverter.get('model', ''))
                ws_inverters.cell(row=row, column=4, value=inverter.get('power', ''))
            except:
                pass
    
    wb.save(output_path)
    return output_path

def fill_memorial_document(data):
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
    template_path = os.path.join(base_dir, "templates", "MODELODEMEMORIALTÉCNICODESCRITIVO.docx")
    output_path = os.path.join(base_dir, "output", "Memorial_Tecnico_Preenchido.docx")

    # Create output directory if it doesn't exist
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    doc = Document(template_path)
    
    # Replace placeholders in the document
    replacements = {
        'XX': str(data.get('total_power', 'XX')),
        '[TENSÃO NOMINAL DA REDE]': data.get('grid_voltage', '[TENSÃO NOMINAL DA REDE]'),
        '[AUTOCONSUMO LOCAL, AUTOCONSUMO REMOTO, GERAÇÃO COMPARTILHADA OU EMUC]': data.get('consumption_type', '[AUTOCONSUMO LOCAL]')
    }
    
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text in paragraph.text:
                paragraph.text = paragraph.text.replace(old_text, new_text)
    
    doc.save(output_path)
    return output_path

def fill_procuracao_document(data):
    base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
    template_path = os.path.join(base_dir, "templates", "ExemploProcuracaoPessoaFisica.docx")
    output_path = os.path.join(base_dir, "output", "Procuracao_Preenchida.docx")

    # Create output directory if it doesn't exist
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    
    doc = Document(template_path)
    
    # Replace placeholders in the document
    replacements = {
        '[NOME DO OUTORGANTE]': data.get('client_name', '[NOME DO OUTORGANTE]'),
        '[CPF DO OUTORGANTE]': data.get('client_cpf', '[CPF DO OUTORGANTE]'),
        '[ENDEREÇO DO OUTORGANTE]': data.get('client_address', '[ENDEREÇO DO OUTORGANTE]')
    }
    
    for paragraph in doc.paragraphs:
        for old_text, new_text in replacements.items():
            if old_text in paragraph.text:
                paragraph.text = paragraph.text.replace(old_text, new_text)
    
    doc.save(output_path)
    return output_path

@automation_bp.route('/calculate-system', methods=['POST'])
@cross_origin()
def calculate_system():
    try:
        data = request.json
        
        modules = data.get('modules', [])
        inverters = data.get('inverters', [])
        
        # Converter strings para números
        for module in modules:
            module['quantity'] = int(module.get('quantity', 0)) if module.get('quantity') else 0
            module['power'] = float(module.get('power', 0)) if module.get('power') else 0
        
        for inverter in inverters:
            inverter['quantity'] = int(inverter.get('quantity', 0)) if inverter.get('quantity') else 0
            inverter['power'] = float(inverter.get('power', 0)) if inverter.get('power') else 0
        
        # Filtrar apenas equipamentos com dados válidos
        valid_modules = [m for m in modules if m['quantity'] > 0 and m['power'] > 0]
        valid_inverters = [i for i in inverters if i['quantity'] > 0 and i['power'] > 0]
        
        if not valid_modules or not valid_inverters:
            return jsonify({'error': 'Dados insuficientes para cálculo'}), 400
        
        # Cálculos básicos
        total_module_power = sum(m['quantity'] * m['power'] for m in valid_modules)
        total_inverter_power = sum(i['quantity'] * i['power'] for i in valid_inverters)
        
        # Cálculos avançados
        cable_calc = calculate_cable_section_advanced(total_inverter_power)
        protection = calculate_protection_devices_advanced(total_inverter_power)
        generation = calculate_energy_generation(valid_modules)
        compatibility = validate_system_compatibility(valid_modules, valid_inverters)
        
        # Análise econômica (estimativa)
        estimated_cost = total_module_power * 3.5  # R$ 3,50/Wp estimado
        economics = calculate_economic_analysis(estimated_cost, generation['monthly_generation_kwh'])
        
        calculations = {
            'power_summary': {
                'total_module_power_kw': round(total_module_power / 1000, 2),
                'total_inverter_power_kw': round(total_inverter_power / 1000, 2),
                'power_ratio': compatibility['power_ratio']
            },
            'generation': {
                'monthly_generation_kwh': generation['monthly_generation_kwh'],
                'annual_generation_kwh': generation['annual_generation_kwh'],
                'system_efficiency': 0.85
            },
            'cable_section_recommendation': f"{cable_calc['recommended_section_mm2']} mm²",
            'voltage_drop_percent': cable_calc['voltage_drop_percent'],
            'protection_devices': {
                'ac_breaker': f"{protection['ac_breaker_a']}A",
                'dc_breaker': f"{protection['dc_breaker_a']}A",
                'dps': f"{protection['dps_class']} - {protection['dps_voltage']}",
                'dr_sensitivity': f"{protection['dr_sensitivity_ma']}mA"
            },
            'compatibility': {
                'status': compatibility['status'],
                'power_ratio': compatibility['power_ratio'],
                'warnings': compatibility['warnings'],
                'errors': compatibility['errors']
            },
            'economic_analysis': {
                'estimated_cost': round(estimated_cost, 2),
                'monthly_savings': economics['monthly_savings_brl'],
                'payback_years': economics['simple_payback_years'],
                'npv': economics['npv_brl'],
                'tir_percent': economics['tir_percent']
            }
        }

        return jsonify(calculations)
        
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@automation_bp.route('/save-form', methods=['POST'])
@cross_origin()
def save_form():
    try:
        data = request.json

        # Criar diretório saved_forms se não existir
        base_dir = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(__file__))))
        saved_forms_dir = os.path.join(base_dir, "saved_forms")
        os.makedirs(saved_forms_dir, exist_ok=True)

        # Gerar nome do arquivo com timestamp
        from datetime import datetime
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        client_name = data.get('client', {}).get('client_name', 'cliente')
        safe_name = ''.join(c for c in client_name if c.isalnum() or c in (' ', '_')).replace(' ', '_')
        filename = f"form_{safe_name}_{timestamp}.json"
        filepath = os.path.join(saved_forms_dir, filename)

        # Salvar JSON
        with open(filepath, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)

        return jsonify({
            'success': True,
            'filename': filename,
            'path': filepath
        })

    except Exception as e:
        import traceback
        error_trace = traceback.format_exc()
        print(f"ERROR in save_form: {error_trace}")
        return jsonify({'error': str(e)}), 500

@automation_bp.route('/health', methods=['GET'])
def health_check():
    return jsonify({'status': 'healthy'})

