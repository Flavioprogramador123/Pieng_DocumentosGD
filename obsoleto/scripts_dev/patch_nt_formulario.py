"""
Atualiza o formulário NT.00020-05 (GUIA 0 e GUIA 1) com tokens revisados.

Gera cópia .xlsx para conferência e atualiza o .xltx oficial.
"""
from __future__ import annotations

import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TEMPLATES = ROOT / 'templates'
XLTX = TEMPLATES / 'NT.00020-05-Anexo-I-Formulario-de-Solicitacao-Grupo-B-templates.xltx'
XLSX_REVISAO = TEMPLATES / 'NT.00020-05-Anexo-I-Formulario-de-Solicitacao-Grupo-B-Tokens-Revisao.xlsx'

# GUIA 0 — linha 7 módulos; linhas 22–31 inversores (1 por linha)
SOLAR_ROW = 7
SOLAR_TOKEN_COLS = {
    'D': '{{POTENCIA_MODULO}}',
    'H': '{{QTD_MODULOS}}',
    'P': '{{AREA_ARRANJO}}',
    'T': '{{FABRICANTE_MODULO}}',
    'AA': '{{MODELO_MODULO}}',
}
PEAK_FORMULA = '=IFERROR((D{row}*H{row})/1000,"")'

INVERTER_FIRST_ROW = 22
INVERTER_LAST_ROW = 31
INVERTER_TOKEN_COLS = {
    'D': '{{FABRICANTE_INVERSOR}}',
    'H': '{{MODELO_INVERSOR}}',
    'L': '{{POTENCIA_INVERSOR}}',
    'P': '{{FAIXA_TENSAO_INVERSOR}}',
    'T': '{{CORRENTE_INVERSOR}}',
    'W': '{{FATOR_POTENCIA}}',
    'Z': '{{RENDIMENTO}}',
    'AC': '{{DHT}}',
}

# GUIA 1 — tokens revisados (sempre célula superior-esquerda do merge)
GUIA1_TOKENS = {
    'AC9': '{{RG_COMPLETO}}',
    'C13': '{{ENDERECO_UC}}',
    'T13': '{{TELEFONE_CELULAR}}',
    'F29': '{{DEMANDA_ALVO_KW}}',
    'P33': '{{COORDENADA_UTM_X}}',
    'Y33': '{{COORDENADA_UTM_Y}}',
}


def _set_token(ws, addr: str, token: str) -> None:
    """Grava token na célula mestre (canto superior-esquerdo se mesclada)."""
    from openpyxl.cell.cell import MergedCell

    cell = ws[addr]
    if isinstance(cell, MergedCell):
        for merged in ws.merged_cells.ranges:
            if addr in merged:
                addr = str(merged).split(':')[0]
                break
    ws[addr] = token


def backup(path: Path) -> Path:
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    dest = path.with_suffix(f'{path.suffix}.bak_{stamp}')
    shutil.copy2(path, dest)
    return dest


def patch_guia0(ws) -> None:
    for col, token in SOLAR_TOKEN_COLS.items():
        ws[f'{col}{SOLAR_ROW}'] = token
    for row in range(SOLAR_ROW, SOLAR_ROW + 10):
        ws[f'K{row}'] = PEAK_FORMULA.format(row=row)

    for row in range(INVERTER_FIRST_ROW, INVERTER_LAST_ROW + 1):
        for col, token in INVERTER_TOKEN_COLS.items():
            ws[f'{col}{row}'] = token


def patch_guia1(ws) -> None:
    for addr, token in GUIA1_TOKENS.items():
        _set_token(ws, addr, token)


def main() -> None:
    if not XLTX.exists():
        raise SystemExit(f'Template não encontrado: {XLTX}')

    backup(XLTX)
    wb = load_workbook(XLTX, keep_vba=False)

    if '0' not in wb.sheetnames or '1' not in wb.sheetnames:
        raise SystemExit(f'Abas esperadas 0/1 não encontradas: {wb.sheetnames}')

    patch_guia0(wb['0'])
    patch_guia1(wb['1'])

    wb.save(XLTX)
    wb.save(XLSX_REVISAO)

    print(f'Atualizado: {XLTX}')
    print(f'Revisão:    {XLSX_REVISAO}')
    print('GUIA 0: K7–K16 fórmula pico; P7 AREA_ARRANJO; inversores linhas 22–31')
    print('GUIA 1: RG_COMPLETO, ENDERECO_UC, TELEFONE_CELULAR, DEMANDA_ALVO_KW, UTM')


if __name__ == '__main__':
    main()
